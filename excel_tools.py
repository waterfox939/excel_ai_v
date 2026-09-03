"""All openpyxl read and write functions, for the CLI's file-based workflows."""
from pathlib import Path
from typing import Any

import openpyxl


def load_workbook(path: str | Path):
    return openpyxl.load_workbook(path)


def list_sheets(path: str | Path) -> list[str]:
    return load_workbook(path).sheetnames


def read_sheet_grid(path: str | Path, sheet_name: str) -> list[list[Any]]:
    """Read an entire sheet as a raw grid of cell values, row by row.

    Use this for free-form sheets (like a financial statement) that aren't
    laid out as a clean table.
    """
    ws = load_workbook(path)[sheet_name]
    return [[cell.value for cell in row] for row in ws.iter_rows()]


def find_header_row(path: str | Path, sheet_name: str, min_non_empty: int = 3) -> int:
    """Return the 1-indexed row number of the first row with at least
    `min_non_empty` non-empty string cells, skipping title rows above it.
    """
    ws = load_workbook(path)[sheet_name]
    for row in ws.iter_rows():
        strings = [c.value for c in row if isinstance(c.value, str) and c.value.strip()]
        if len(strings) >= min_non_empty:
            return row[0].row
    raise ValueError(f"No header row found in sheet {sheet_name!r}")


def read_table(
    path: str | Path,
    sheet_name: str,
    header_row: int | None = None,
    forward_fill_columns: list[str] | None = None,
) -> list[dict[str, Any]]:
    """Read a sheet into a list of row dicts keyed by header text.

    header_row: 1-indexed row containing column headers; auto-detected if omitted.
    forward_fill_columns: header names whose blank cells should inherit the
    value from the row above — used for grouped columns like order number,
    where the value is only written on the first line of a multi-line order.
    """
    ws = load_workbook(path)[sheet_name]
    header_row = header_row or find_header_row(path, sheet_name)
    headers = [c.value for c in ws[header_row]]

    data_start = header_row + 1
    next_row_values = [c.value for c in ws[data_start]]
    is_header_continuation = any(
        v is not None for v in next_row_values
    ) and all(
        v is None or (isinstance(v, str) and headers[i]) for i, v in enumerate(next_row_values)
    )
    if is_header_continuation:
        headers = [f"{h}/{v}" if v else h for h, v in zip(headers, next_row_values)]
        data_start += 1

    rows = []
    last_values: dict[str, Any] = {}
    for row in ws.iter_rows(min_row=data_start):
        record = {}
        for header, cell in zip(headers, row):
            if not header:
                continue
            value = cell.value
            if value is None and forward_fill_columns and header in forward_fill_columns:
                value = last_values.get(header)
            elif header in (forward_fill_columns or []):
                last_values[header] = value
            record[header] = value
        if any(v is not None for v in record.values()):
            rows.append(record)
    return rows


def get_cell(path: str | Path, sheet_name: str, row: int, col: int) -> Any:
    return load_workbook(path)[sheet_name].cell(row=row, column=col).value


def set_cells(path: str | Path, sheet_name: str, updates: list[tuple[int, int, Any]]) -> None:
    """Write multiple cells and save once. updates: list of (row, col, value)."""
    wb = load_workbook(path)
    ws = wb[sheet_name]
    for row, col, value in updates:
        ws.cell(row=row, column=col, value=value)
    wb.save(path)


def append_row(path: str | Path, sheet_name: str, values: list[Any]) -> int:
    """Append a row after the last used row. Returns the new row number."""
    wb = load_workbook(path)
    ws = wb[sheet_name]
    new_row = ws.max_row + 1
    for col, value in enumerate(values, start=1):
        ws.cell(row=new_row, column=col, value=value)
    wb.save(path)
    return new_row
